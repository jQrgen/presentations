# Blockchain Game Jam pitch deck (python-pptx, 16:9). Two variants from one script:
#   python make_deck.py ntnu     -> /workspace/gamejam-deck/ntnu-blockchain-game-jam.pptx
#   python make_deck.py general  -> /workspace/gamejam-deck/general/blockchain-game-jam.pptx  (Organizer-neutral)
# Style adapted from the earlier Nexa deck helpers (dark theme, gold accents). Reconstructed 2026-09-30 after the
# original build dirs were removed from the box.
import re, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

VARIANT = sys.argv[1] if len(sys.argv) > 1 else "ntnu"
G = VARIANT == "general"
def V(ntnu, general): return general if G else ntnu

BG = RGBColor(0x0D, 0x10, 0x17); PANEL = RGBColor(0x17, 0x1C, 0x26); PANEL2 = RGBColor(0x20, 0x26, 0x33)
GOLD = RGBColor(0xFF, 0xD2, 0x3F); WHITE = RGBColor(0xF2, 0xF4, 0xF8); MUTED = RGBColor(0x9A, 0xA4, 0xB5)
GREEN = RGBColor(0x2E, 0xC4, 0x8A); ORANGE = RGBColor(0xFF, 0x8A, 0x3D); BLUE = RGBColor(0x4F, 0xA3, 0xFF)
PURPLE = RGBColor(0x9B, 0x5C, 0xD6); PINK = RGBColor(0xFF, 0x5C, 0xA8); RED = RGBColor(0xE5, 0x48, 0x4D)
FONT = "Inter"
A = "/workspace/gamejam-deck/build/assets/"
OUT = V("/workspace/gamejam-deck/ntnu-blockchain-game-jam.pptx", "/workspace/gamejam-deck/general/blockchain-game-jam.pptx")
VIDEO_URL = "https://www.youtube.com/watch?v=NHAPE0JRhCM"

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]; N = [0]
STAGES = ["RENTED", "OWNED", "PORTABLE", "PROOF", "THE JAM", "TOOLS", "WIN-WIN", "THE ASK"]

def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, radius=None, lw=1.5):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if radius is not None: sh.adjustments[0] = radius
    return sh

def text(s, x, y, w, h, content, size=22, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, font=FONT):
    """Any [bracketed] text becomes a pink placeholder run."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.line_spacing = spacing
        runs = p if isinstance(p, list) else [(p, {})]
        split = []
        for r in runs:
            if isinstance(r, str): r = (r, {})
            t0, o0 = r
            for piece in re.split(r'(\[[^\]]*\])', t0):
                if piece: split.append((piece, dict(o0, color=PINK) if piece.startswith("[") else o0))
        for t, o in split:
            run = para.add_run(); run.text = t; f = run.font
            f.name = o.get("font", font); f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False); f.color.rgb = o.get("color", color)
    return tb

def shape_text(sh, content, size=18, color=WHITE, bold=False, align=PP_ALIGN.CENTER):
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = content if isinstance(content, list) else [content]
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        runs = l if isinstance(l, list) else [(l, {})]
        for r in runs:
            if isinstance(r, str): r = (r, {})
            t, o = r
            run = p.add_run(); run.text = t; run.font.name = FONT; run.font.size = Pt(o.get("size", size))
            run.font.bold = o.get("bold", bold); run.font.color.rgb = o.get("color", color)

def notes(s, t): s.notes_slide.notes_text_frame.text = t.strip()

def tracker(s, stage):
    w, gap, x0 = 1.33, 0.08, 0.9
    for i, st in enumerate(STAGES):
        cur = st == stage; passed = STAGES.index(stage) > i
        fill = GOLD if cur else (PANEL2 if passed else PANEL)
        sh = rect(s, x0 + i * (w + gap), 6.55, w, 0.28, fill, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        shape_text(sh, [[(st, {"bold": True, "size": 9, "color": BG if cur else (WHITE if passed else MUTED)})]])
        sh.text_frame.word_wrap = False

def new_slide(title=None, kicker=None, stage=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    N[0] += 1
    if title:
        rect(s, 0.6, 0.55, 0.12, 0.95, GOLD)
        if kicker: text(s, 0.9, 0.42, 10, 0.4, kicker.upper(), size=14, color=GOLD, bold=True)
        text(s, 0.9, 0.72, 11.6, 0.9, title, size=36, bold=True)
    if stage: tracker(s, stage)
    text(s, 0.6, 7.0, 9, 0.3, V("NTNU Blockchain Game Jam", "Blockchain Game Jam") + "  ·  proposal  ·  jQrgen  ·  Bitcoin Unlimited", size=11, color=MUTED)
    text(s, 12.0, 7.0, 0.75, 0.3, str(N[0]), size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    return s

def arrow(s, x1, y1, x2, y2, col=GOLD, width=2.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = col; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))

def rbox(s, x, y, w, h, lines, fill=PANEL, line_col=None, size=17, radius=0.12):
    sh = rect(s, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, line=line_col, radius=radius)
    shape_text(sh, lines, size=size); return sh

# ---- the running example: ONE card (the "Defense Boost 100" equip card from the Nexa Warriors demo, 29:25)
def game_card(s, x, y, sc=1.0, state="game", caption=None):
    """state: game (locked, orange), wallet (gold, owned), free (green, in a new game)"""
    col = {"game": ORANGE, "wallet": GOLD, "free": GREEN}[state]
    w, h = 1.7 * sc, 2.4 * sc
    rect(s, x, y, w, h, PANEL2, MSO_SHAPE.ROUNDED_RECTANGLE, line=col, radius=0.08, lw=2.5 * sc)
    text(s, x, y + 0.1 * sc, w, 0.3 * sc, "EQUIP", size=max(8, int(10 * sc)), bold=True, color=col, align=PP_ALIGN.CENTER)
    rect(s, x + w / 2 - 0.42 * sc, y + 0.48 * sc, 0.84 * sc, 0.95 * sc, col, MSO_SHAPE.PLAQUE)
    text(s, x + 0.05, y + 1.5 * sc, w - 0.1, 0.35 * sc, "DEFENSE", size=max(8, int(13 * sc)), bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.05, y + 1.8 * sc, w - 0.1, 0.35 * sc, "BOOST 100", size=max(8, int(13 * sc)), bold=True, align=PP_ALIGN.CENTER)
    if caption:
        text(s, x - 0.5, y + h + 0.08, w + 1.0, 0.4, caption, size=13, color=col, bold=True, align=PP_ALIGN.CENTER)
    return w, h

def cage(s, x, y, w, h, label):
    rect(s, x, y, w, h, None, MSO_SHAPE.ROUNDED_RECTANGLE, line=ORANGE, radius=0.06, lw=2)
    for i in range(1, 6):
        xx = x + i * w / 6
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xx), Inches(y), Inches(xx), Inches(y + h))
        c.line.color.rgb = ORANGE; c.line.width = Pt(1.2)
    text(s, x, y - 0.42, w, 0.4, label, size=14, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)

# =============================================================== 1 TITLE
s = new_slide()
s.shapes.add_picture(A + "nexa-logo.png", Inches(0.9), Inches(0.8), height=Inches(0.6))
text(s, 0.9, 1.9, 8, 0.5, V("A PROPOSAL FOR NTNU TRONDHEIM", "A PROPOSAL FOR UNIVERSITIES & STUDENT ASSOCIATIONS"), size=16, bold=True, color=GOLD)
text(s, 0.9, 2.35, 8, 1.0, "Break it out", size=54, bold=True)
text(s, 0.9, 3.2, 8, 1.0, "of the game", size=54, bold=True, color=GOLD)
text(s, 0.9, 4.35, 8, 0.5, "A 48-hour blockchain game jam for students", size=22)
rect(s, 0.95, 5.0, 1.2, 0.06, GOLD)
text(s, 0.9, 5.2, 8, 0.45, "jQrgen  ·  Bitcoin Unlimited / Nexa", size=18, color=MUTED)
text(s, 0.9, 5.6, 8, 0.45, V("[date TBD]  ·  [venue TBD]  ·  [partner student organisation TBD]",
                             "[Organizer: university, student association...]  ·  [date TBD]  ·  [venue TBD]"), size=15)
s.shapes.add_picture(A + "bu-logo.png", Inches(0.95), Inches(6.25), height=Inches(0.38))
game_card(s, 9.3, 1.7, sc=1.35, state="wallet")
text(s, 8.3, 5.15, 4.3, 0.5, "Who owns this card?", size=24, bold=True, align=PP_ALIGN.CENTER)
text(s, 5.9, 6.2, 6.85, 0.6, [[("Based on: ", {"bold": True, "color": GOLD}), ("\u201cMoving Beyond Centralized Gaming using the Nexa Blockchain\u201d", {"italic": True})],
     [("with Steven van den Meiracker  \u00b7  Nexa on YouTube, 2026", {})]], size=11, color=MUTED, align=PP_ALIGN.RIGHT)
notes(s, """
[STAGE: hold up ONE printed playing card - the "Defense Boost 100" equip card from the Nexa Warriors demo. Say nothing for 3 seconds.]

HOOK: "Someone earned this card. Tonight's question is simple: who owns it - the player, or the game?" [PAUSE]

- Introduce yourself: jQrgen, Bitcoin Unlimited, working on the Nexa blockchain. """ + V("[Optional: NTNU alumnus? - confirm before saying it.]", "[Optional: your connection to the [Organizer] - add if relevant.]") + """
- One-sentence pitch: a 48-hour game jam where """ + V("NTNU students build", "students at [Organizer] build") + """ games whose items belong to the players and can leave the game.
- This card is our running example. It shows up on almost every slide: locked in a game, then in a wallet, then in a brand-new game built at the jam.
- Pink text = a placeholder we still need to fill in together """ + V("(date, venue, partner, budget)", "(Organizer, date, venue, budget)") + """.

The 'Based on' line on the slide credits the source video; the last slide repeats it with a QR code.
Source for the whole deck: "Moving Beyond Centralized Gaming using the Nexa Blockchain: with Steven van den Meiracker", Nexa YouTube channel, uploaded 10 Jul 2026 (https://www.youtube.com/watch?v=NHAPE0JRhCM). Webinar hosted by Danielle at EvolvH3R. The card "Defense Boost 100" appears in the demo at 00:29:25.
""")

# =============================================================== 2 PROBLEM: you rent
s = new_slide("You don't own your games. You rent them.", "The problem", "RENTED")
cage(s, 1.1, 2.35, 3.6, 3.1, "THE GAME'S SERVER")
game_card(s, 2.05, 2.7, sc=1.0, state="game")
items = [("Steam", "you buy a licence to play, and that licence can be revoked", "01:35"),
         ("Switch 2", "the game card is \"just the licence key\"", "01:35"),
         ("Your sword", "hours of upgrades live on closed infrastructure", "08:47")]
for i, (h, b, t) in enumerate(items):
    y = 2.3 + i * 1.12
    rect(s, 5.6, y, 6.9, 0.95, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, 5.85, y + 0.12, 1.9, 0.7, h, size=20, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.75, y + 0.12, 4.6, 0.7, b, size=16, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 5.85, 11.6, 0.5, "Publishers can modify, remove or discontinue your items, without you really knowing.", size=17, color=MUTED)
notes(s, """
[STAGE: put the printed card behind your back / in a closed box. "Right now, this card lives here - on somebody else's server."]

- Steven van den Meiracker (Senior Development Consultant & Solution Architect, Bitcoin Unlimited) opens with ownership: "you're really renting your games. You're not really owning them." On Steam you buy "a license to play the game and that license can be revoked" [01:35].
- His own anecdote: he bought a Switch 2 game and the physical card was "just the license key" [01:35-02:06].
- The reward problem: the sword you upgraded "through hours and hours of effort", the Pokemon you caught - "all exist within the game on closed infrastructure" [08:17-08:47].
- Game assets "live on the server of the game"; developers and publishers "can modify or remove those assets, can discontinue an asset... without the players even really knowing about it" [18:04-18:35].
- Ask the room: "Who here has lost items when a game shut down or banned an account?" [WAIT for hands.]
""")

# =============================================================== 3 HONEST FRAMING
s = new_slide("Blockchain won't \"revolutionize gaming\"", "Honest framing from the talk", "RENTED")
text(s, 0.9, 1.85, 11.6, 1.6, [[("\u201cYou can't just put the word blockchain onto something and say, well yeah, we made it better.\u201d", {"italic": True})]],
     size=28, color=WHITE, spacing=1.15)
text(s, 0.9, 3.35, 11, 0.4, "Steven van den Meiracker, Bitcoin Unlimited  ·  03:06", size=14, color=MUTED)
rbox(s, 0.9, 4.2, 5.5, 1.75, [[("Games are thriving", {"bold": True, "size": 21, "color": GOLD})],
                              [("Studios do well. Modding communities are strong.", {"size": 16})]])
rbox(s, 7.0, 4.2, 5.5, 1.75, [[("So blockchain must earn its place", {"bold": True, "size": 21, "color": GREEN})],
                              [("Solve a real problem, or unlock something new.", {"size": 16})]])
notes(s, """
[STAGE: lower your voice. This slide builds trust with a sceptical student audience.]

- The talk is explicitly NOT "blockchain fixes gaming": "it's not really going to revolutionize gaming. You can't just bolt blockchain onto something" [01:04]. Studios "are thriving", and there is a healthy modding ecosystem - he mentions Cities: Skylines mods that get promoted and become officially supported [02:36].
- The method: "look at a problem and look at the benefits of what blockchain can bring and then determine how that best solves a specific problem or unlocks something new" [03:06-03:37].
- He repeats this in Q&A: "It has to make things better. It has to make it easier... more cost-effective" [42:56].
- The host notes the Web2 gaming community was hostile to crypto ("scams") [40:52-41:23]. Say it out loud: "If you're sceptical - good. That's the attitude we want at the jam."
- This becomes the jam's first judging criterion (slide 12).
""")

# =============================================================== 4 LANGUAGE OF OWNERSHIP
s = new_slide("A blockchain is a language of ownership", "The key insight", "OWNED")
for i, (h, b, c) in enumerate([("Distributed", "no central server to switch off", BLUE),
                               ("Trustless", "strangers trade, no middleman", PURPLE),
                               ("Immutable", "the record can't be rewritten", GREEN)]):
    x = 0.9 + i * 2.55
    rect(s, x, 2.1, 2.35, 2.3, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    rect(s, x, 2.1, 2.35, 0.1, c)
    text(s, x + 0.15, 2.4, 2.05, 0.5, h, size=21, bold=True, color=c, align=PP_ALIGN.CENTER)
    text(s, x + 0.15, 3.05, 2.05, 1.2, b, size=16, align=PP_ALIGN.CENTER)
arrow(s, 8.65, 3.25, 9.35, 3.25)
game_card(s, 9.75, 1.9, sc=1.0, state="wallet", caption="lives in YOUR wallet")
text(s, 0.9, 5.1, 11.6, 1.1, [[("The game issues and uses the card. ", {}), ("The player holds it.", {"bold": True, "color": GOLD})]],
     size=28)
notes(s, """
[STAGE: take the card out from behind your back and hold it up high. "This is the moment it becomes yours."]

- The three pillars as the talk lists them: distributed (the Nexa network "runs on hundreds of nodes around the world" - no single point of failure) [03:37-04:07]; trustless ("two complete strangers on opposite sides of the world can interact in a transaction" without a middleman) [04:38-05:09]; immutable ("once something is written... it cannot be changed and it can be proven that it has not changed") [05:09].
- Put together: "the blockchain really is a language of ownership" - Steven credits the phrase to Andrew Stone, BU's head developer [06:12].
- For games: minted on-chain assets "actually live in your wallet. They don't live inside the game. They're accessed by the game. They're used by the game. They're issued by the game" [19:37].
- Note: in-game economies are not new (Second Life, EVE Online, WoW auction house) [16:31-17:34]. Putting them on a chain alone "is not adding anything new" [17:34]. That's the setup for the next slide.
""")

# =============================================================== 5 PORTABLE - the new dimension
s = new_slide("The new part: items can leave the game", "What's actually new", "PORTABLE")
cage(s, 0.9, 2.45, 2.4, 2.9, "ORIGINAL GAME")
game_card(s, 1.45, 2.7, sc=0.76, state="game")
arrow(s, 3.45, 3.9, 4.55, 3.9)
game_card(s, 4.7, 2.45, sc=1.0, state="wallet", caption="player's wallet")
targets = [("Battle arena", 2.0), ("Trading tool", 3.05), ("Companion app", 4.1), ("Mod in another game", 5.15)]
for lab, y in targets:
    arrow(s, 6.55, 3.65, 8.3, y + 0.35, col=GREEN, width=2)
    rbox(s, 8.4, y, 4.1, 0.72, [[(lab, {"bold": True, "size": 17, "color": GREEN})]], fill=PANEL)
text(s, 0.9, 5.95, 11.6, 0.5, "\u2026built by anyone, \u201cwithout asking permission.\u201d  (21:10)", size=18, color=MUTED)
notes(s, """
[STAGE: walk the card physically from one side of the stage to the other while you talk - out of the "game" corner, into your pocket (the wallet), then hand it to a volunteer in the front row: "and now it's in HER game."]

- Steven: putting achievements on chain "is a pretty easy step. There's not really anything different in that" [09:19]. The new idea is to couple it with a "true ownership license... You truly do own this asset both in the game and outside the game" [09:19-09:50].
- With that licence, you "could actually build that entirely outside the game without requiring the permission of the gaming developer" - an add-on game, a battling arena, a different trading model, a swap mechanism [19:37-20:07].
- "Assets flow out from a game and assets flow in from other games or other experiences if the developer wants to support those assets" [20:39]. People "can build tournaments around it, other trading tools, companion apps, all without asking permission" [21:10].
- Cross-game example: a Nexa Warriors card "could be supported in a mod that a community member builds in a whole other experience", e.g. a champion tournament game [21:10-21:42].
- Why a studio would want this: popular assets outside the game create "more demand for those assets... the only way you can get them is back in that original game" [21:42-22:13].
- His own caveat: "will it be a successful model? Only time will tell... you have to give it a go" [22:44]. => A game jam is exactly "giving it a go", 20 times in one weekend.
""")

# =============================================================== 6 PROOF: Nexa Warriors
s = new_slide("It already exists: Nexa Warriors", "Proof, not a promise", "PROOF")
stats = [("165", "cards, each an NFT"), ("6", "rarity tiers, set at minting"), ("32", "cards per deck, 8 can be locked")]
for i, (n, l) in enumerate(stats):
    y = 1.95 + i * 1.3
    rect(s, 0.9, y, 5.6, 1.1, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, 1.05, y + 0.1, 1.5, 0.9, n, size=40, bold=True, color=GOLD, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    text(s, 2.65, y + 0.1, 3.75, 0.9, l, size=17, anchor=MSO_ANCHOR.MIDDLE)
feats = [("NFTs", "the cards"), ("Tokens", "shards: burn + upgrade a card"), ("Smart contracts", "winner takes one card"),
         ("Wally Wallet", "scan a QR to log in"), ("On-chain check", "game verifies you own it")]
for i, (h, b) in enumerate(feats):
    y = 1.95 + i * 0.76
    rect(s, 7.0, y, 5.5, 0.64, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
    text(s, 7.2, y + 0.08, 2.1, 0.5, h, size=16, bold=True, color=GREEN, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.3, y + 0.08, 3.1, 0.5, b, size=15, anchor=MSO_ANCHOR.MIDDLE)
sh = rect(s, 0.9, 5.95, 11.6, 0.45, None, MSO_SHAPE.ROUNDED_RECTANGLE, line=ORANGE, radius=0.3)
shape_text(sh, [[("In development, testnet only at the time of the talk. No launch date announced yet.", {"size": 14, "color": ORANGE, "bold": True})]])
notes(s, """
[STAGE: if you can, show the 1-minute demo clip from the video (00:26:20 onwards), or the marketplace on your laptop. Otherwise just point at the card: "this is one of the 165."]

- Nexa Warriors: a turn-based electronic trading card game "built natively on the Nexa network" [23:14-23:44]. 165 cards, every card minted as an NFT; six rarity tiers "enforced at the minting level"; cards come from starter decks, booster packs, the open market and gameplay rewards [23:14-24:16].
- Tokens: you earn "shards" and burn a lower-tier card plus shards to produce a higher-tier card [24:47].
- Smart contracts: in a two-player game your cards are locked into a contract and "the winner gets one of the opponent cards"; of your 32 cards you can lock 8 out of risk [24:47-25:49].
- Wally Wallet identity: "you scan the game with your wallet", and the game "is validating on chain that I do own those assets" [24:16, 28:54-29:25].
- Demo moment: Steven lists a duplicate "Defense Boost 100" for 100,000 NEX as a signed partial transaction - on testnet, no real value [29:25-29:57]. That's our running-example card.
- Proceeds from packs and decks go to Bitcoin Unlimited, a not-for-profit, to fund development [25:49].
- Status (be honest): "This is all currently only on our test net... These will be brought across to mainnet when we officially launch" [26:50]. Public testnet play not encouraged yet [31:34-32:06]. Launch "definitely this year... hopefully within a few months" [47:32]. The Nexa forum June 2026 dev update: "the date is not set yet" (forum.nexa.org/t/1721); the August 2026 update still lists Nexa Warriors as "on the way" (forum.nexa.org/t/1751).
""")

# =============================================================== 7 THE JAM
s = new_slide(V("The idea: NTNU Blockchain Game Jam", "The idea: a blockchain game jam"), "The format", "THE JAM")
blocks = [("48h", "one weekend"), ("[2-4]", "students per team"), ("1", "theme"), ("[N]", "teams")]
for i, (n, l) in enumerate(blocks):
    x = 0.9 + i * 2.95
    rect(s, x, 1.95, 2.7, 1.65, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, x, 2.05, 2.7, 0.9, n, size=44, bold=True, color=PINK if n.startswith("[") else GOLD, align=PP_ALIGN.CENTER)
    text(s, x, 2.95, 2.7, 0.5, l, size=16, color=MUTED, align=PP_ALIGN.CENTER)
rect(s, 0.9, 3.95, 11.6, 2.25, PANEL2, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, line=GOLD)
text(s, 1.2, 4.1, 5, 0.45, "THEME", size=14, bold=True, color=GOLD)
text(s, 1.2, 4.5, 8.4, 0.8, "\u201cBreak it out of the game\u201d", size=34, bold=True)
text(s, 1.2, 5.3, 8.4, 0.8, "Build a game where the items belong to the players, and can live on outside it.", size=17, color=MUTED)
game_card(s, 10.55, 4.12, sc=0.78, state="free")
notes(s, """
[STAGE: slow down, one block at a time. Then read the theme like a movie title.]

- The talk does not mention a game jam, students or """ + V("NTNU", "any specific organizer") + """ - this is our proposal, built on the talk's central idea (assets that "flow out from a game" [20:39]) and its own conclusion that you "have to give it a go" [22:44].
- 48 hours is the classic game-jam length; the exact format, team size [2-4 TBD] and number of teams [N TBD] are for us to agree with """ + V("[partner student organisation TBD]", "the [Organizer] (university, student association, etc.)") + """.
- Theme is deliberately one line. Teams may build: a new game that uses on-chain items; an add-on (arena, tournament, trading tool, companion app) - the exact examples Steven lists [19:37-21:10].
- Optional bonus track [only if BU confirms assets/testnet access]: build something that accepts Nexa Warriors cards, like the "champion tournament game" he describes [21:10-21:42].
- The card on the slide is now green: at the jam, it gets a second life in a new game.
""")

# =============================================================== 8 WHY NEXA
s = new_slide("Why build it on Nexa?", "Why blockchain, why this one", "THE JAM")
why = [("Near-zero fees", "\u201ca fraction of a fraction of a cent\u201d", "15:29", GREEN),
       ("Native tokens", "minted and moved like the base coin", "15:29", GOLD),
       ("Smart contracts", "lock a card into a match, enforce the prize", "15:29 · 25:17", BLUE),
       ("Partial transactions", "\u201cI have A, I want B\u201d offers anyone can fill", "15:59", PURPLE),
       ("Wally identity", "log in with your wallet, no password", "24:16", ORANGE)]
for i, (h, b, t, c) in enumerate(why):
    y = 1.9 + i * 0.86
    rect(s, 0.9, y, 11.6, 0.74, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
    rect(s, 0.9, y + 0.14, 0.07, 0.46, c)
    text(s, 1.15, y + 0.1, 3.3, 0.55, h, size=19, bold=True, color=c, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.5, y + 0.1, 6.6, 0.55, b, size=17, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 11.0, y + 0.1, 1.4, 0.55, t, size=11, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
notes(s, """
[STAGE: tap the card once per row - "fees: this card costs nothing to move. Tokens: ..." etc.]

- Nexa as Steven describes it: "Bitcoin version three" - same UTXO model and tokenomics as Bitcoin, proof-of-work, fair launch, halvings; he claims "60,000 transactions a second plus compared to seven" and "near zero fees" [14:25-14:56]. (Throughput figure is the speaker's claim - I have not re-verified it for this deck; say "designed for" rather than "does".)
- Fees "a fraction of a fraction of a cent"; native tokens that are "part of consensus" and transacted the same way as the core currency; "programmable with smart contracts" [15:29-15:59].
- Open-outcry partial transactions: "I have A and I'd like B... someone at some point could pick that up... and execute the transaction" [15:59-16:31] - Steven used exactly this to list his card for sale [29:25-29:57]. Perfect for a jam team building a trading tool.
- Wally Wallet identity: scanning a QR code logs you into the game and shares your assets [24:16, 28:24-28:54].
- Why it matters for a jam: players in a 48h demo won't wait for or pay high fees; cheap, fast transactions make on-chain game actions practical.
""")

# =============================================================== 9 TOOLS
s = new_slide("What teams build with", "The toolkit", "TOOLS")
tools = [("Build On Nexa (BON)", "docs: mint NFTs, create tokens, use Wally", "build.nexa.org"),
         ("Wally Wallet", "players' wallet + login by QR", "wallywallet.org"),
         ("Nexa AI skills", "skill files that teach coding agents Nexa", "gitlab.com/nexa/nexaaiskills"),
         ("Nexa testnet", "free test coins, real rules, no real money", "[confirm faucet/testnet access]")]
for i, (h, b, u) in enumerate(tools):
    x = 0.9 + (i % 2) * 5.95; y = 1.9 + (i // 2) * 2.0
    rect(s, x, y, 5.65, 1.75, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    rect(s, x, y + 0.25, 0.07, 1.25, GOLD)
    text(s, x + 0.3, y + 0.18, 5.2, 0.5, h, size=22, bold=True, color=GOLD)
    text(s, x + 0.3, y + 0.72, 5.2, 0.5, b, size=16)
    text(s, x + 0.3, y + 1.18, 5.2, 0.4, u, size=13, color=MUTED)
text(s, 0.9, 5.95, 11.6, 0.5, [[("Mentors on site: ", {"bold": True}), ("[Bitcoin Unlimited / Nexa developers TBD]", {})]], size=16)
notes(s, """
[STAGE: hold up your phone with Wally Wallet open - "every player at the jam leaves with this on their phone."]

- Build On Nexa (BON): Steven: "how to mint NFTs, how to create tokens, how to interact with Wally wallet, all of that material is becoming available in our... build on Nexa program" [39:21-39:51] (auto-captions say "bond"). BON launched publicly Oct 2025 at build.nexa.org with TypeScript/JavaScript and Kotlin support; BU also runs BONFires workshops and BONQuests developer challenges (Nexa forum posts 1587 and 1650; nexa.org/build lists libnexa-ts, NexScript and more). BONQuests may be a natural umbrella for the jam - [confirm with BU].
- Wally Wallet (wallywallet.org): the wallet used in the demo; the cards appear in Wally [28:24].
- Nexa AI skills (gitlab.com/nexa/nexaaiskills): a corpus of skills that AI coding agents (Claude, Cursor and similar) load so they "produce correct, idiomatic, and secure Nexa application code on the first attempt"; covers tokens and groups, wallet connection, NPL smart contracts, transactions and more (repo README). Not mentioned in the video - this is jQrgen's addition. For a 48h jam this is a big speed-up.
- Testnet: Nexa Warriors itself runs on testnet [26:50]. Confirm faucet access and whether the jam may use testnet before promising it.
- Mentors: [names TBD].
""")

# =============================================================== 10 SCHEDULE (running example through the weekend)
s = new_slide("48 hours, one card, many new games", "Schedule", "THE JAM")
steps = [("FRI 18:00", "Kickoff + tournament", "[if Nexa Warriors is ready]", "wallet"),
         ("FRI 20:00", "Teams form, theme revealed", "", "wallet"),
         ("SAT", "Build: mint, code, playtest", "mentors on site", "free"),
         ("SUN 14:00", "Demos: show the card in YOUR game", "", "free"),
         ("SUN 16:00", "Judging + prizes", "[prizes TBD]", "free")]
for i, (t, h, sub, st) in enumerate(steps):
    x = 0.9 + i * 2.37
    rect(s, x, 1.95, 2.17, 3.9, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    text(s, x, 2.08, 2.17, 0.4, t, size=15, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    game_card(s, x + 0.66, 2.6, sc=0.5, state=st)
    text(s, x + 0.1, 3.95, 1.97, 1.1, h, size=16, bold=True, align=PP_ALIGN.CENTER)
    if sub: text(s, x + 0.1, 5.05, 1.97, 0.7, sub, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 4: arrow(s, x + 2.19, 3.2, x + 2.35, 3.2, width=2)
text(s, 0.9, 5.95, 11.6, 0.5, "Times are a sketch.  [date TBD]  ·  [venue TBD]  ·  [food + overnight access TBD]", size=15, color=MUTED)
notes(s, """
[STAGE: move the printed card along the five columns with your finger as you talk.]

- This timeline is a proposal, not from the video. All times are a sketch to agree with """ + V("[partner student organisation TBD]", "the [Organizer]") + """.
- Kickoff mini-tournament: the host of the webinar floated "a learning education tournament" once the game is ready [48:02]. Nexa Warriors has a free trial mode "without owning any cards" [32:36] and "randomly generates your assets" [46:00], which makes it a zero-cost icebreaker - IF it has launched or BU gives testnet access. Otherwise, replace with a 20-minute live "mint your first token" demo from the BON docs.
- Friday: theme revealed, teams formed (mix programmers, artists, designers).
- Saturday: build. Mentors [TBD] on site.
- Sunday: demos - every team must show an item leaving (or entering) their game. Judging and prizes [TBD].
""")

# =============================================================== 11 WIN-WIN
s = new_slide("Everyone leaves with something", "What each side gets", "WIN-WIN")
cols = [("Students", GOLD, ["Ship a real game in 48h", "Hands-on blockchain", "Portfolio + [prizes TBD]", "Mentors from Nexa devs"]),
        V(("NTNU + [partner]", BLUE, ["A new kind of jam on campus", "Cost: [budget TBD]", "Industry contact", "[other benefits TBD]"]),
          ("[Organizer]", BLUE, ["A new kind of jam for members", "Cost: [budget TBD]", "Industry contact", "[other benefits TBD]"])),
        ("BU / Nexa", GREEN, ["New builders on Nexa", "Feedback on BON docs", "Ideas for portable items", V("Talent in Trondheim", "Meet student talent")])]
for i, (h, c, its) in enumerate(cols):
    x = 0.9 + i * 3.95
    rect(s, x, 1.9, 3.7, 4.35, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(s, x, 1.9, 3.7, 0.1, c)
    text(s, x + 0.2, 2.1, 3.3, 0.5, h, size=20, bold=True, color=c)
    for j, it in enumerate(its):
        text(s, x + 0.2, 2.9 + j * 0.78, 3.3, 0.7, [[("\u25B8  ", {"color": c, "bold": True}), (it, {})]], size=16)
notes(s, """
[STAGE: point at three different parts of the room - "you, the university, and us."]

- Students: a shipped game and portfolio piece; hands-on experience with NFTs, tokens, smart contracts and wallets - the building blocks Steven demoed [24:47-25:17]; prizes [TBD].
""" + V("""- NTNU / [partner student organisation TBD]: NTNU is Norway's largest university, ~43,500 students, headquartered in Trondheim (85% study there) - ntnu.edu/about and ntnu.edu/facts. A jam fits its science-and-technology profile. Concrete benefits beyond that are for the partner to define [TBD] - do not promise anything not agreed.""",
"""- What the Organizer gets ([Organizer]: university, faculty, student association, hackathon club, etc.): a hands-on, low-cost event for its students or members, and direct contact with protocol developers. Add organizer-specific facts (size, profile, previous events) here once known [TBD] - verify them and cite a source. Concrete benefits are for the Organizer to define [TBD] - do not promise anything not agreed.""") + """
- BU / Nexa: Steven's stated goal for Nexa Warriors is "an onboarding path to bring people into the space" [44:27-44:58] and he invites people to build via BON [39:21]. A jam delivers builders plus concrete feedback on the docs (the Nexa AI skills repo even keeps a gap log, SKILL-UPDATES.md, for exactly this).
- Do not claim any sponsorship, funding or partnership that is not yet agreed.
""")

# =============================================================== 12 JUDGING
s = new_slide("How we judge: the talk's own test", "Judging criteria", "WIN-WIN")
crit = [("1", "Does blockchain earn its place?", "solves a problem or unlocks something new", "03:06"),
        ("2", "Does an item leave (or enter) the game?", "true ownership, not just a database", "20:39"),
        ("3", "Could a Web2 gamer play it?", "easy wallet onboarding, free first try", "44:27"),
        ("4", "Is it fun?", "effort vs. reward balance", "08:17")]
for i, (n, h, b, t) in enumerate(crit):
    y = 1.85 + i * 1.02
    rect(s, 0.9, y, 11.6, 0.9, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    sh = rect(s, 1.1, y + 0.16, 0.6, 0.6, GOLD, MSO_SHAPE.OVAL)
    shape_text(sh, [[(n, {"bold": True, "size": 20, "color": BG})]])
    text(s, 1.95, y + 0.08, 6.0, 0.45, h, size=19, bold=True)
    text(s, 1.95, y + 0.48, 6.5, 0.4, b, size=14, color=MUTED)
    text(s, 10.8, y + 0.2, 1.5, 0.5, t, size=12, color=MUTED, align=PP_ALIGN.RIGHT)
text(s, 0.9, 5.98, 11.6, 0.4, "Jury: [TBD]", size=15, color=MUTED)
notes(s, """
[STAGE: count on your fingers, 1 to 4.]

1. From the opening: you have to "look at a problem... and then determine how that best solves a specific problem or unlocks something new" [03:06-03:37]; repeated in Q&A: "unless it ticks these boxes, then it's not really going to be successful" [42:56-43:26].
2. The core idea: assets "flow out from a game and... flow in from other games or other experiences" [20:39].
3. Onboarding: the Q&A asks whether Web2 players can play without understanding wallets; Steven: "you need to get a wallet. You need to fund that wallet" and they aim to make it "as easy as possible"; the free-game model lets you play with an unfunded wallet [43:57-46:00].
4. Fun: "If it's too hard to achieve anything, then you give up. If it's too easy, then the reward's not meaningful" [08:17].
- Jury composition [TBD] - e.g. """ + V("NTNU staff, student org, BU developer", "Organizer staff or student-association board, BU developer") + """; to be agreed.
""")

# =============================================================== 13 BUDGET & ASK
s = new_slide("What we need to make it happen", "Budget & ask", "THE ASK")
COST = V("[NOK TBD]", "[cost TBD]")
rows = [("Venue for 48h", "[venue TBD]", COST), ("Food + drinks", "[participants TBD]", COST),
        ("Prizes", "[prize structure TBD]", COST), ("Mentors + travel", "[BU staff TBD]", COST),
        V(("Marketing on campus", "[partner org channels]", COST), ("Marketing", "[Organizer channels]", COST))]
rect(s, 0.9, 1.9, 7.2, 4.35, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
for i, (a, b, c) in enumerate(rows):
    y = 2.1 + i * 0.8
    text(s, 1.15, y, 2.8, 0.5, a, size=17, bold=True)
    text(s, 3.95, y, 2.6, 0.5, b, size=14)
    text(s, 6.45, y, 1.5, 0.5, c, size=14, align=PP_ALIGN.RIGHT)
    if i < 4: rect(s, 1.15, y + 0.62, 6.7, 0.02, PANEL2)
rect(s, 8.5, 1.9, 4.0, 4.35, PANEL2, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, line=GOLD)
text(s, 8.75, 2.1, 3.5, 0.45, "THE ASK", size=14, bold=True, color=GOLD)
for j, (h, b) in enumerate([(V("From NTNU / [partner]", "From the [Organizer]"), "room, promotion, jury seat"), ("From BU / Nexa", "mentors, [funding TBD]"), ("Sponsors", "[sponsors TBD]")]):
    text(s, 8.75, 2.65 + j * 1.15, 3.5, 0.45, h, size=17, bold=True)
    text(s, 8.75, 3.08 + j * 1.15, 3.5, 0.5, b, size=15, color=MUTED)
notes(s, """
[STAGE: stop moving. Stand still for the ask - this is the business part.]

- Every number on this slide is a placeholder. Nothing here comes from the video; no sponsor, budget or prize has been agreed.
- To fill in: venue [TBD], participant cap [TBD], food, prize structure and amounts [TBD], BU mentor travel [TBD], marketing via """ + V("[partner student organisation TBD] channels.", "[Organizer] channels. Currency [TBD].") + """
- The ask is split: """ + V("NTNU / partner provides", "the Organizer provides") + """ room, promotion and a jury seat; BU / Nexa provides mentors and possibly funding [confirm internally - do not promise]; external sponsors [TBD].
- If asked about prizes paid in NEX: keep it a placeholder until BU agrees; mention that the demo marketplace was on testnet with no real value [26:50].
""")

# =============================================================== 14 CTA
s = new_slide()
game_card(s, 1.55, 1.3, sc=1.2, state="free")
# watch-the-talk reference with QR code
s.shapes.add_picture(A + "qr-talk.png", Inches(1.1), Inches(4.6), height=Inches(1.3))
text(s, 2.55, 4.6, 2.35, 0.4, "WATCH THE TALK", size=13, bold=True, color=GOLD)
text(s, 2.55, 4.95, 2.35, 1.0, [["\u201cMoving Beyond Centralized Gaming using the Nexa Blockchain\u201d"], ["Steven van den Meiracker"], ["Nexa on YouTube, 2026"]], size=10, spacing=1.05)
text(s, 1.05, 6.0, 3.9, 0.35, "youtube.com/watch?v=NHAPE0JRhCM", size=11, color=MUTED)
text(s, 5.0, 1.7, 7.6, 0.5, "SO, WHO OWNS THIS CARD?", size=18, bold=True, color=GOLD)
text(s, 5.0, 2.2, 7.6, 1.8, [[("The player.", {})], V([("And the next game? ", {}), ("NTNU.", {"color": GOLD})], [("Next game: ", {}), ("[Organizer]", {})])], size=38, bold=True, spacing=1.05)
steps = [V("Pick a date and a partner student organisation", "Pick a date with the [Organizer]"), "Confirm venue, mentors and budget", V("Open sign-ups on campus", "Open sign-ups")]
for i, st in enumerate(steps):
    y = 4.3 + i * 0.62
    sh = rect(s, 5.0, y, 0.42, 0.42, GOLD, MSO_SHAPE.OVAL)
    shape_text(sh, [[(str(i + 1), {"bold": True, "size": 14, "color": BG})]])
    text(s, 5.6, y + 0.0, 7.0, 0.45, st, size=18, anchor=MSO_ANCHOR.MIDDLE)
text(s, 5.0, 6.2, 7.6, 0.4, "jQrgen  ·  [contact email TBD]  ·  build.nexa.org", size=15, color=MUTED)
notes(s, """
[STAGE: walk to the front row, hand the printed card to the """ + V("NTNU contact / student-org lead", "Organizer's contact (university / student-association lead)") + """. "This one is yours now. Let's build the next game around it."]

- Watch the talk (QR code on the slide): "Moving Beyond Centralized Gaming using the Nexa Blockchain - with Steven van den Meiracker", Nexa on YouTube, 2026 - https://www.youtube.com/watch?v=NHAPE0JRhCM. Point people to the Nexa Warriors demo from 00:23:14.
- Close the loop from slide 1: "Who owns this card?" - the player. And at the jam, students build the next game around it.
- Three next steps: (1) """ + V("agree a date and a partner student organisation [TBD]", "agree a date with the [Organizer]") + """; (2) confirm venue, mentors and budget [TBD]; (3) open sign-ups.
- Leave your contact: [contact email TBD]. Point to build.nexa.org for anyone who wants to start before the jam; Nexa X and Discord were the channels named in the talk [39:21].
- Quote to end on (Steven, 38:51): "It really has to bring something different... provide a new innovative angle or a new way to do things."
""")

prs.save(OUT); print("saved", OUT, N[0], "slides")
